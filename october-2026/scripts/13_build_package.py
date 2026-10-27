# -*- coding: utf-8 -*-
# Build the submission package for Computational Brain & Behavior (v4) from the draft sources:
# assemble the manuscript (named, with the title page bound in front, and anonymous), the title
# page, the cover letter and Online Resource 1; resolve every number token from the frozen v3
# result files (11_resolve.py); convert with pandoc to DOCX (TeX math becomes Word equations);
# add automatic page numbers; clean the document metadata; export PDFs with Microsoft Word; copy
# the figures; and write word and page counts to draft/build/build_stats.json.
# v4 changes against v3: newline escapes restored (the v3 file did not parse), output goes to
# package/, an anonymous manuscript is built next to the named one, math is checked for OMML,
# tables use a compact 8 pt style, the supplement is named ESM_1 as the guidelines ask.
# Usage: python 13_build_package.py
import os, re, sys, json, shutil, subprocess, zipfile, tempfile, datetime
import importlib.util

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(HERE)
DRAFT = os.path.join(BASE, "draft")
BUILD = os.path.join(DRAFT, "build")
PKG = os.path.join(BASE, "package")
os.makedirs(BUILD, exist_ok=True)
os.makedirs(PKG, exist_ok=True)
spec = importlib.util.spec_from_file_location("resolve", os.path.join(HERE, "11_resolve.py"))
RES = importlib.util.module_from_spec(spec)
spec.loader.exec_module(RES)
AUTHOR = "Ismayil Zeynalov"
REPO = "https://github.com/ismayilzeynal/forgetting-curve-minimum-evidence"
TITLE = open(os.path.join(DRAFT, "title.txt"), encoding="utf-8").read().strip()
RUNHEAD = open(os.path.join(DRAFT, "runninghead.txt"), encoding="utf-8").read().strip()
_today = datetime.date.today()
DATE = os.environ.get("CBB_DATE") or "%d %s %d" % (_today.day, _today.strftime("%B"), _today.year)
TMP = tempfile.mkdtemp(prefix="cbb_build_")

# r1 fix (desk-r1-01, methods-r1-01): the code availability statement says only what is public.
# Default: the public repository holds the July 2026 bundle (plan, confirmatory code and results);
# the October code, results and addendum (repo-v4/) are available on request. After repo-v4/ has been
# pushed and checked against repo-v4/MANIFEST.sha256, build with --repo-v4-public to name the
# repository for everything.
REPO_V4_PUBLIC = "--repo-v4-public" in sys.argv
if REPO_V4_PUBLIC:
    CODE_NAMED = ("The analysis code that reproduces every reported number from the public traces, the result "
                  "files, the dated analysis plan of 12 July 2026 and the exploratory addendum of 14 October "
                  "2026 are available at " + REPO + ".")
    CODE_ANON = ("The analysis code that reproduces every reported number from the public traces, the result "
                 "files, the dated analysis plan of 12 July 2026 and the exploratory addendum of 14 October "
                 "2026 are available in a public repository; the link is given on the title page.")
else:
    CODE_NAMED = ("The dated analysis plan of 12 July 2026, with the code and results of the confirmatory run, is "
                  "public at " + REPO + ". The code that reproduces every number reported here from the public "
                  "traces, the result files and the exploratory addendum of 14 October 2026 are available from "
                  "the author on request.")
    CODE_ANON = ("The dated analysis plan of 12 July 2026, with the code and results of the confirmatory run, is "
                 "in a public repository whose link is given on the title page. The code that reproduces every "
                 "number reported here from the public traces, the result files and the exploratory addendum of "
                 "14 October 2026 are available from the author on request.")
# r2 fix (desk-r2-09): the cover letter says plainly what is public and what is on request
if REPO_V4_PUBLIC:
    CODE_COVER = ("The code and result files of every analysis reported here, the plan and the addendum are public "
                  "in the repository named in the code availability statement.")
else:
    CODE_COVER = ("The link to that plan and code is given in the code availability statement; the code and result "
                  "files of the analyses reported here and the addendum are available from the author on request.")
VARIANTS = {
    "named": {
        "AUTHOR_CONTRIB": "Ismayil Zeynalov is the sole author and is responsible for the conception and design "
                          "of the study, the analyses and the manuscript.",
        "CODE_AVAIL": CODE_NAMED,
    },
    "anon": {
        "AUTHOR_CONTRIB": "The sole author is responsible for the conception and design of the study, the analyses "
                          "and the manuscript.",
        "CODE_AVAIL": CODE_ANON,
    },
}


def rd(name):
    return open(os.path.join(DRAFT, name), encoding="utf-8").read().replace("\r\n", "\n")


def fill(text, extra=None):
    text = text.replace("{{TITLE}}", TITLE).replace("{{RUNNING_HEAD}}", RUNHEAD).replace("{{DATE}}", DATE)
    for k, v in (extra or {}).items():
        text = text.replace("{{%s}}" % k, v)
    return text


def assemble():
    """Write the token sources of every document to draft/build/."""
    figs = {}
    for blk in rd("figures.md").split("@@")[1:]:
        name, body = blk.split("\n", 1)
        figs[name.strip()] = body.strip()
    parts = ["part0_front.md", "part1_intro_background_theory.md", "part2_methods.md", "part3_results_a.md",
             "part3_results_b.md", "part4_discussion.md", "part5_back.md", "part9_references.md"]
    body = "\n\n".join(rd(f).strip() for f in parts) + "\n"
    body = body.replace("{{TABLE2}}", rd("table2.md").strip()).replace("{{TABLE3}}", rd("table3.md").strip())
    for k, v in figs.items():
        body = body.replace("{{%s}}" % k, v)
    out = {}
    for var, extra in VARIANTS.items():
        out["manuscript_" + var] = fill(body, extra)
    out["supplement"] = fill(rd("supplement_body.md").replace("{{TABLES5}}", rd("tableS5.md").strip()))
    out["cover_letter"] = fill(rd("cover_letter_src.md"), {"CODE_COVER": CODE_COVER})
    out["title_page"] = fill(rd("title_page_src.md"), {"CODE_AVAIL": VARIANTS["named"]["CODE_AVAIL"]})
    paths = {}
    for k, v in out.items():
        p = os.path.join(BUILD, k + "_src.md")
        open(p, "w", encoding="utf-8", newline="\n").write(v)
        paths[k] = p
    return paths


def resolve(src):
    dst = src.replace("_src.md", "_resolved.md")
    RES.main(src, dst, src.replace("_src.md", "_numbers.tsv"))
    return dst


def reference_docx():
    ref = os.path.join(TMP, "ref.docx")
    with open(ref, "wb") as fh:
        fh.write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                                capture_output=True, check=True).stdout)
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_LINE_SPACING
    d = Document(ref)
    for st in d.styles:
        try:
            if st.type == 1:
                st.font.name = "Times New Roman"
                st.font.color.rgb = None
        except Exception:
            pass
    names = [s.name for s in d.styles]
    for nm in ("Normal", "Body Text", "First Paragraph", "Abstract", "Block Text"):
        if nm in names:
            s = d.styles[nm]
            s.font.size = Pt(12)
            s.font.name = "Times New Roman"
            s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            s.paragraph_format.space_after = Pt(6)
    if "Compact" in names:          # table cells
        s = d.styles["Compact"]
        s.font.size = Pt(8)
        s.font.name = "Times New Roman"
        s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        s.paragraph_format.space_after = Pt(0)
        s.paragraph_format.space_before = Pt(0)
    for nm, sz in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12.5), ("Heading 3", 12)):
        if nm in names:
            s = d.styles[nm]
            s.font.size = Pt(sz)
            s.font.name = "Times New Roman"
            s.font.bold = True
            s.font.italic = False
            s.font.color.rgb = RGBColor(0, 0, 0)
    # theme font attributes override the explicit font name, so drop them from every style
    from docx.oxml.ns import qn
    for st in d.styles:
        rpr = st.element.find(qn("w:rPr"))
        rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if rf is not None:
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin = sec.bottom_margin = Cm(2.5)
    d.save(ref)
    return ref


def pandoc(md, out, ref, hard_breaks=False):
    fmt = "markdown" + ("+hard_line_breaks" if hard_breaks else "")
    subprocess.run(["pandoc", md, "-f", fmt, "-t", "docx", "-o", out, "--reference-doc", ref,
                    "--resource-path", DRAFT], check=True)


def add_page_numbers(path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = Document(path)
    for sec in d.sections:
        ft = sec.footer
        ft.is_linked_to_previous = False
        p = ft.paragraphs[0] if ft.paragraphs else ft.add_paragraph()
        p.alignment = 1
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)
    d.save(path)


def clean_meta(path, title, author):
    from docx import Document
    d = Document(path)
    cp = d.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.title = title
    cp.subject = ""
    cp.comments = ""
    cp.keywords = ""
    cp.category = ""
    cp.content_status = ""
    cp.identifier = ""
    cp.language = "en-US"
    cp.version = ""
    cp.revision = 1
    d.save(path)
    zin = zipfile.ZipFile(path)
    items = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    if "docProps/app.xml" in items:
        x = items["docProps/app.xml"].decode("utf-8")
        for tag in ("Company", "Manager", "HyperlinkBase"):
            x = re.sub(r"<%s>.*?</%s>" % (tag, tag), "", x, flags=re.S)
        x = re.sub(r"<Template>.*?</Template>", "<Template>Normal.dotm</Template>", x, flags=re.S)
        # r2 fix (desk-r2-13): drop the template's stale statistics (all optional elements); Word
        # recomputes them on save, and the build records Word's own counts in build_stats.json
        for tag in ("TotalTime", "Pages", "Words", "Characters", "CharactersWithSpaces", "Lines", "Paragraphs"):
            x = re.sub(r"<%s>.*?</%s>|<%s\s*/>" % (tag, tag, tag), "", x, flags=re.S)
        items["docProps/app.xml"] = x.encode("utf-8")
    items.pop("docProps/custom.xml", None)
    if "[Content_Types].xml" in items:
        items["[Content_Types].xml"] = re.sub(rb'<Override PartName="/docProps/custom.xml"[^>]*/>', b"",
                                              items["[Content_Types].xml"])
    if "_rels/.rels" in items:
        items["_rels/.rels"] = re.sub(rb'<Relationship[^>]*docProps/custom.xml"[^>]*/>', b"", items["_rels/.rels"])
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, b in items.items():
            zout.writestr(n, b)
    os.replace(tmp, path)


def omml_count(path):
    with zipfile.ZipFile(path) as z:
        x = z.read("word/document.xml").decode("utf-8")
    return len(re.findall(r"<m:oMath[ >]", x)), len(re.findall(r"<m:oMathPara[ >]", x))


def to_pdf(word, docx, pdf):
    doc = word.Documents.Open(os.path.abspath(docx), ReadOnly=True, AddToRecentFiles=False)
    try:
        pages = doc.ComputeStatistics(2)
        words_ = doc.ComputeStatistics(0)
        doc.ExportAsFixedFormat(os.path.abspath(pdf), 17, False, 0, 0, 1, 1, 0, True, True, 0, True, True, False)
    finally:
        doc.Close(False)
    return pages, words_


def plain_words(text):
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)(\{[^}]*\})?", " ", text)
    text = re.sub(r"\$\$.*?\$\$", " ", text, flags=re.S)
    text = re.sub(r"\$[^$\n]+\$", " X ", text)
    text = re.sub(r"<[^>]+>", " ", text)
    return len(re.findall(r"[A-Za-z0-9À-ɏ−][A-Za-z0-9À-ɏ'.,%/+−-]*", text))


def body_words(md, start, stops):
    t = open(md, encoding="utf-8").read()
    i = t.find(start)
    j = min([t.find(p, i + 1) for p in stops if t.find(p, i + 1) > 0] + [len(t)])
    keep = []
    for para in t[i:j].split("\n\n"):
        s = para.strip()
        if s.startswith("|") or s.startswith("**Fig.") or s.startswith("**Table") or s.startswith("*Note.*") \
                or s.startswith("!["):
            continue
        keep.append(s)
    return plain_words("\n\n".join(keep))


def main():
    paths = assemble()
    resolved = {k: resolve(p) for k, p in paths.items()}
    ref = reference_docx()
    bound = os.path.join(BUILD, "manuscript_with_title_page_resolved.md")
    pb = '\n\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'
    open(bound, "w", encoding="utf-8", newline="\n").write(
        open(resolved["title_page"], encoding="utf-8").read() + pb + open(resolved["manuscript_named"], encoding="utf-8").read())
    targets = [
        (bound, "Manuscript_with_title_page.docx", TITLE, AUTHOR, False),
        (resolved["manuscript_anon"], "Manuscript_anonymous.docx", TITLE, "", False),
        (resolved["title_page"], "Title_page.docx", TITLE + " - title page", AUTHOR, False),
        (resolved["cover_letter"], "Cover_letter.docx", "Cover letter", AUTHOR, True),
        (resolved["supplement"], "ESM_1.docx", TITLE + " - Online Resource 1", AUTHOR, False),
    ]
    stats = dict(build_date=DATE, files={})
    for md, name, ttl, auth, hb in targets:
        p = os.path.join(PKG, name)
        pandoc(md, p, ref, hard_breaks=hb)
        add_page_numbers(p)
        clean_meta(p, ttl, auth)
        stats["files"][name] = dict(omml_inline_and_display=omml_count(p)[0], omml_display=omml_count(p)[1])
    import win32com.client, pythoncom
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        for name in [t[1] for t in targets]:
            pg, wd = to_pdf(word, os.path.join(PKG, name), os.path.join(PKG, name.replace(".docx", ".pdf")))
            stats["files"][name].update(word_pages=pg, word_words=wd)
    finally:
        word.Quit()
    for i in range(1, 7):
        for ext in ("png", "pdf"):
            shutil.copy2(os.path.join(DRAFT, "figures", "Fig%d.%s" % (i, ext)), os.path.join(PKG, "Fig%d.%s" % (i, ext)))
    mt = open(resolved["manuscript_named"], encoding="utf-8").read()
    abstract = mt.split("## Abstract", 1)[1].split("**Keywords:**", 1)[0].strip()
    keywords = mt.split("**Keywords:**", 1)[1].split("\n", 1)[0].strip()
    open(os.path.join(PKG, "abstract.txt"), "w", encoding="utf-8", newline="\n").write(abstract + "\n")
    open(os.path.join(PKG, "keywords.txt"), "w", encoding="utf-8", newline="\n").write(
        "\n".join(k.strip() for k in keywords.split(";")) + "\n")
    refs = mt.split("## References", 1)[1]
    stats.update(
        abstract_words=len(abstract.split()),
        keywords=len(keywords.split(";")),
        main_text_words=body_words(resolved["manuscript_named"], "## 1 Introduction",
                                   ["## Supplementary Information", "## Statements and Declarations"]),
        references=len(re.findall(r"^\S.*\(\d{4}[a-z]?(?:, [A-Z][a-z]+ \d{1,2})?\)\.", refs, flags=re.M)),
        code_statement="repo-v4 public" if REPO_V4_PUBLIC else "July bundle public, rest on request",
        tables_main=len(re.findall(r"^\*\*Table \d+\.\*\*", mt, flags=re.M)),
        figures_main=len(re.findall(r"^\*\*Fig\. \d+\*\*", mt, flags=re.M)),
        supplement_tables=len(re.findall(r"^\*\*Table S\d+\.\*\*", open(resolved["supplement"], encoding="utf-8").read(), flags=re.M)),
        cover_letter_words=plain_words(open(resolved["cover_letter"], encoding="utf-8").read()),
        tokens={k: sum(1 for _ in open(p.replace("_src.md", "_numbers.tsv"), encoding="utf-8")) - 1 for k, p in paths.items()},
    )
    json.dump(stats, open(os.path.join(BUILD, "build_stats.json"), "w", encoding="utf-8"), indent=1)
    print(json.dumps(stats, indent=1))
    shutil.rmtree(TMP, ignore_errors=True)


if __name__ == "__main__":
    main()
